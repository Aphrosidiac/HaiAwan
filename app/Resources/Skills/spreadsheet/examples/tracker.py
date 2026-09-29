"""A person-facing shortlist in openpyxl: one sheet, best row first and highlighted, live formulas,
real number formats, frozen header, filters. Run: python3 tracker.py output/spreadsheet/shortlist.xlsx
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

rows = [
    # name, monthly price (RM), setup fee (RM), seats, rating, link
    ("Option A", 149, 0, 5, 4.6, "https://example.com/a"),
    ("Option B", 99, 300, 3, 4.2, "https://example.com/b"),
    ("Option C", 199, 0, 10, 4.8, "https://example.com/c"),
]
rows.sort(key=lambda r: r[4], reverse=True)  # best option first (here: highest rating)
headers = ["Option", "Monthly (RM)", "Setup (RM)", "Seats", "Rating", "First-year cost (RM)", "Link"]

wb = Workbook()
ws = wb.active
ws.title = "Shortlist"
ws.append(headers)
for r in rows:
    ws.append(list(r[:5]) + [None, r[5]])

last = ws.max_row
for i in range(2, last + 1):
    ws[f"F{i}"] = f"=B{i}*12+C{i}"          # a formula, not a pasted number
    ws[f"G{i}"].hyperlink = ws[f"G{i}"].value
    ws[f"G{i}"].style = "Hyperlink"

head_fill = PatternFill("solid", fgColor="EDEBE4")
for c in ws[1]:
    c.font = Font(bold=True)
    c.fill = head_fill
    c.alignment = Alignment(vertical="center")
for col in ("B", "C", "F"):
    for cell in ws[col][1:]:
        cell.number_format = '"RM" #,##0'
for cell in ws["E"][1:]:
    cell.number_format = "0.0"

# Highlight the best option (row 2 after sorting).
best = PatternFill("solid", fgColor="E8F7B8")
for c in ws[2]:
    c.fill = best

widths = [14, 14, 12, 8, 8, 20, 28]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"

wb.save(sys.argv[1] if len(sys.argv) > 1 else "shortlist.xlsx")
